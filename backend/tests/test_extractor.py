"""Unit tests for Module 1: Dual-Stream Document Ingestion (backend/tests/test_extractor.py)."""

from __future__ import annotations

import io
from unittest.mock import patch

import fitz
from docx import Document
from docx.shared import Pt
from lxml import etree

from backend.models.schemas import BoundingBoxBlock
from backend.services.extractor import (
    detect_column_interleaving,
    extract_document,
    extract_docx_xml,
    extract_layout_blocks,
    extract_linear_stream,
    xy_cut_sort,
)


def _create_single_column_pdf_with_header_footer() -> bytes:
    """Create a synthetic PDF with header (< 72pt), body content, and footer (> height - 72pt)."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)  # 8.5 x 11 inches

    # Header inside top 1 inch (y = 40pt < 72pt)
    page.insert_text((72, 40), "sarah.chen@example.com | (415) 555-0142", fontsize=10, fontname="helv")

    # Main body content (y between 100pt and 650pt)
    page.insert_text((72, 110), "Sarah Chen", fontsize=15, fontname="helv")
    page.insert_text((72, 140), "Work Experience", fontsize=14, fontname="helv")
    page.insert_text(
        (72, 165),
        "Senior Product Manager | Zendex Inc. | Mar 2022 - Present",
        fontsize=11,
        fontname="helv",
    )
    page.insert_text(
        (72, 185),
        "- Led roadmap for enterprise analytics suite, $8M ARR in year one",
        fontsize=11,
        fontname="helv",
    )

    # Footer inside bottom 1 inch (y = 760pt > 792 - 72 = 720pt)
    page.insert_text((72, 760), "Page 1 of 1 - Confidential Resume", fontsize=9, fontname="helv")

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def _create_two_column_pdf() -> bytes:
    """Create a synthetic two-column PDF to test Stream A interleaving vs Stream B XY-Cut++."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)

    # Left Column (x = 72) and Right Column (x = 340) on the same vertical baselines
    page.insert_text((72, 120), "Work Experience", fontsize=14, fontname="helv")
    page.insert_text((340, 120), "Skills", fontsize=14, fontname="helv")

    page.insert_text((72, 150), "Senior Python Developer", fontsize=11, fontname="helv")
    page.insert_text((340, 150), "Python, Django, AWS", fontsize=11, fontname="helv")

    page.insert_text((72, 180), "Built scalable APIs", fontsize=11, fontname="helv")
    page.insert_text((340, 180), "Docker, PostgreSQL", fontsize=11, fontname="helv")

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def _create_blank_scanned_pdf() -> bytes:
    """Create a PDF with < 50 characters of native text to trigger Stage 0 OCR fallback."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 120), "Scan", fontsize=11, fontname="helv")
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def test_stage0_ocr_fallback_and_degradation_penalty():
    pdf_bytes = _create_blank_scanned_pdf()

    with patch("backend.services.extractor.pytesseract.image_to_string") as mock_ocr:
        mock_ocr.return_value = "Sarah Chen\nsarah.chen@example.com\nWork Experience\nSenior Engineer"
        res = extract_document(pdf_bytes, filename="scan.pdf", ocr_enabled=True)
        assert res.ocr_triggered is True
        assert res.confidence_multiplier == 0.85
        assert "Sarah Chen" in res.linear_stream_text

    # When OCR is disabled by platform profile (e.g., standard Greenhouse without OCR)
    res_no_ocr = extract_document(pdf_bytes, filename="scan.pdf", ocr_enabled=False)
    assert res_no_ocr.ocr_triggered is True
    assert res_no_ocr.linear_stream_text == ""
    assert res_no_ocr.is_unreadable_or_encrypted is True


def test_stream_a_header_footer_clipping():
    pdf_bytes = _create_single_column_pdf_with_header_footer()
    linear_text, dropped = extract_linear_stream(pdf_bytes, clip_headers_footers=True)

    assert any("sarah.chen@example.com" in d for d in dropped)
    assert any("Page 1 of 1" in d for d in dropped)
    assert "sarah.chen@example.com" not in linear_text
    assert "Senior Product Manager" in linear_text


def test_stream_b_xy_cut_multi_column_and_interleaving_detector():
    pdf_bytes = _create_two_column_pdf()
    res = extract_document(pdf_bytes, filename="two_col.pdf")

    # Stream B (XY-Cut++) must read Column 1 top-to-bottom before Column 2 top-to-bottom
    layout_lines = [ln.strip() for ln in res.layout_sorted_text.splitlines() if ln.strip()]
    col1_idx = layout_lines.index("Built scalable APIs")
    col2_header_idx = layout_lines.index("Skills")
    assert col1_idx < col2_header_idx

    # Column Interleaving Detector must flag lines concatenated across columns in Stream A
    assert len(res.interleaved_lines) >= 1
    assert any(
        flag.left_block_text == "Work Experience" and flag.right_block_text == "Skills"
        for flag in res.interleaved_lines
    )


def test_xy_cut_preserves_narrow_key_value_gaps():
    """Verify narrow horizontal gaps (< 35pt) on the same baseline remain intact for Key : Value tables."""
    pdf_bytes = _create_single_column_pdf_with_header_footer()
    _, blocks, _, _, _, _ = extract_layout_blocks(pdf_bytes)
    assert len(blocks) >= 3

    # Construct synthetic Key : Value blocks separated by 20pt (< 35pt threshold)
    kv_blocks = [
        BoundingBoxBlock(x0=72.0, y0=200.0, x1=140.0, y1=212.0, text="Languages:", block_no=0, page_no=0),
        BoundingBoxBlock(x0=160.0, y0=200.0, x1=280.0, y1=212.0, text="Python, SQL", block_no=1, page_no=0),
    ]
    sorted_kv = xy_cut_sort(kv_blocks, min_x_gap=35.0)
    assert len(sorted_kv) == 2
    assert sorted_kv[0].text == "Languages:"
    assert sorted_kv[1].text == "Python, SQL"


def test_docx_xml_inspector_tables_textboxes_and_fonts():
    doc = Document()
    # Add header contact text
    section = doc.sections[0]
    section.header.paragraphs[0].text = "header.contact@example.com"

    # Add normal paragraph with unapproved font (Comic Sans MS)
    p = doc.add_paragraph()
    run = p.add_run("Sarah Chen - Resume Body")
    run.font.name = "Comic Sans MS"
    run.font.size = Pt(11)

    # Add a table (<w:tbl>)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Python"

    # Inject a floating text box (<w:txbxContent>) into paragraph XML
    txbx_xml = etree.fromstring(
        '<w:txbxContent xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:p><w:r><w:t>Floating Sidebar Text</w:t></w:r></w:p>"
        "</w:txbxContent>"
    )
    p._p.append(txbx_xml)

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    res = extract_docx_xml(docx_bytes)
    assert res.file_format == "docx"
    assert res.docx_audit is not None
    assert res.docx_audit.has_tables is True
    assert res.docx_audit.has_text_boxes is True
    assert "Comic Sans MS" in res.docx_audit.unapproved_fonts
    assert "header.contact@example.com" in res.dropped_header_footer_text
    assert res.has_tables_or_textboxes is True
